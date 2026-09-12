#!/usr/bin/env python3
"""Rebuild editable RO-018 business documents and PDFs from scoped Markdown sources."""
from __future__ import annotations

import argparse
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[2]
SOURCES = ROOT / 'research/product-engineering/aliyun-phase1-2026-09-07'
JOBS = [('01-partner-recharge.md', '阿里云第一期充值申请说明'),
        ('02-compute-negotiation.md', '阿里云算力规划与商务谈判方案')]


def font(run, size: float = 10, bold: bool = False) -> None:
    run.font.name = 'Noto Sans CJK SC'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    run.font.size = Pt(size)
    run.bold = bold


def inline(paragraph, value: str, size: float = 10) -> None:
    """Support the two inline constructs used by these sources without raw Markdown leakage."""
    parts = re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(https://[^)]+\))', value)
    for part in parts:
        match = re.fullmatch(r'\[([^\]]+)\]\((https://[^)]+)\)', part)
        if match:
            relationship = paragraph.part.relate_to(match[2], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), relationship)
            run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '146B70'); props.append(color)
            fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC'); props.append(fonts)
            text = OxmlElement('w:t'); text.text = match[1]
            run.append(props); run.append(text); link.append(run); paragraph._p.append(link)
        else:
            bold = part.startswith('**') and part.endswith('**')
            font(paragraph.add_run(part[2:-2] if bold else part), size, bold)


def table(document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    grid = document.add_table(rows=0, cols=len(rows[0]))
    grid.style = 'Table Grid'; grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    available = 17.2
    widths = {2: [3.5, 13.7], 3: [4.0, 4.3, 8.9], 4: [3.7, 2.5, 2.7, 8.3]}.get(len(rows[0]))
    if len(rows[0]) == 4 and '单台候选规格' in rows[0]: widths = [4.3, 1.4, 4.4, 7.1]
    if len(rows[0]) == 4 and '三个月' in rows[0]: widths = [4.5, 2.4, 2.4, 7.9]
    if rows[0][0] == '参数': widths = [3.5, 6.85, 6.85]
    if rows[0][0] == '节点用途': widths = [6.5, 5.2, 5.5]
    for col, column in enumerate(grid.columns):
        column.width = Cm(widths[col] if widths else available / len(rows[0]))
    for index, row in enumerate(rows):
        cells = grid.add_row().cells
        tr_pr = cells[0]._tc.getparent().get_or_add_trPr(); tr_pr.append(OxmlElement('w:cantSplit'))
        if index == 0: tr_pr.append(OxmlElement('w:tblHeader'))
        for col, value in enumerate(row):
            cells[col].width = Cm(widths[col] if widths else available / len(row))
            paragraph = cells[col].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2.5)
            paragraph.paragraph_format.space_before = Pt(2.5)
            paragraph.paragraph_format.line_spacing = 1.08
            inline(paragraph, value, 8.7)
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9F0EF'); cells[col]._tc.get_or_add_tcPr().append(shade)
                for run in paragraph.runs: run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build(source: Path, destination: Path, title: str) -> None:
    document = Document(); section = document.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.9); section.right_margin = Cm(1.9)
    section.header_distance = Cm(.65); section.footer_distance = Cm(.65)
    document.core_properties.title = title
    document.core_properties.author = '上海生声不息科技有限公司'
    document.core_properties.subject = '第一期云资源规划｜2026年9月7日｜待审批/征询'
    normal = document.styles['Normal']; normal.font.name = 'Noto Sans CJK SC'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.1
    for name, size in [('Title', 23), ('Heading 1', 16), ('Heading 2', 12)]:
        style = document.styles[name]; style.font.name = 'Noto Sans CJK SC'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('174D50')
        style.paragraph_format.space_before = Pt(6); style.paragraph_format.space_after = Pt(5)
    head = section.header.paragraphs[0]; font(head.add_run('生声不息  /  '+title), 8)
    foot = section.footer.paragraphs[0]
    font(foot.add_run('2026年9月7日  ·  ' + ('内部审批稿' if '充值' in title else '需求征询稿') + ' '*8 + '第 '), 8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    font(foot.add_run(' 页'), 8)
    lines = source.read_text().splitlines(); cursor = 0
    while cursor < len(lines):
        line = lines[cursor].strip(); cursor += 1
        if not line: continue
        if line == '---': document.add_page_break(); continue
        if line.startswith('|'):
            group = [line]
            while cursor < len(lines) and lines[cursor].strip().startswith('|'):
                group.append(lines[cursor].strip()); cursor += 1
            table(document, group); continue
        if line.startswith('# '):
            paragraph = document.add_paragraph(style='Title'); inline(paragraph, line[2:], 23)
        elif line.startswith('## '):
            paragraph = document.add_paragraph(style='Heading 1'); inline(paragraph, line[3:], 16)
        elif line.startswith('### '):
            paragraph = document.add_paragraph(style='Heading 2'); inline(paragraph, line[4:], 12)
        else:
            paragraph = document.add_paragraph(); inline(paragraph, line)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out-dir', type=Path, default=ROOT / 'research/product-engineering/aliyun-phase1-2026-09-07/deliverables')
    args = parser.parse_args(); args.out_dir.mkdir(parents=True, exist_ok=True)
    for source, title in JOBS:
        path = args.out_dir / (title + '.docx'); build(SOURCES / source, path, title)
        subprocess.run(['libreoffice', '-env:UserInstallation=file:///tmp/vox-ro018-lo', '--headless', '--convert-to', 'pdf', '--outdir', str(args.out_dir), str(path)], check=True, cwd=ROOT)
        if not path.with_suffix('.pdf').is_file(): raise RuntimeError(f'PDF missing: {path}')
        print(path)


if __name__ == '__main__':
    main()
